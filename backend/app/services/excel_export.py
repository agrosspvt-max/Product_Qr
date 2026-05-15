"""Excel export for a generation batch."""
from __future__ import annotations

import os
from datetime import datetime
from typing import Iterable

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import settings
from app.models.qr_code import QRCode
from app.utils.logger import get_logger

logger = get_logger(__name__)


COLUMNS = [
    "Code",
    "Product",
    "Quantity",
    "Manufacturing",
    "Short Link",
    "WhatsApp Number",
    "Zoho Sync Status",
]


def _month_name(month: int) -> str:
    return [
        "Jan", "Feb", "Mar", "Apr", "May", "Jun",
        "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
    ][month - 1]


def build_dataframe(codes: Iterable[QRCode]) -> pd.DataFrame:
    rows = []
    for c in codes:
        product_name = c.product.name if c.product else ""
        quantity_label = c.quantity.label if c.quantity else ""
        wa_number = c.whatsapp_preset.whatsapp_number if c.whatsapp_preset else ""
        manufacturing = f"{_month_name(c.manufacturing_month)} {c.manufacturing_year}"
        rows.append({
            "Code": c.code,
            "Product": product_name,
            "Quantity": quantity_label,
            "Manufacturing": manufacturing,
            "Short Link": c.short_link,
            "WhatsApp Number": wa_number,
            "Zoho Sync Status": c.zoho_sync_status,
        })
    return pd.DataFrame(rows, columns=COLUMNS)


def export_batch_to_excel(batch_name: str, codes: Iterable[QRCode]) -> str:
    """Write an Excel file for the batch; return the absolute path."""
    os.makedirs(settings.EXCEL_EXPORT_DIR, exist_ok=True)

    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    filename = f"batch_{batch_name}_{ts}.xlsx"
    path = os.path.abspath(os.path.join(settings.EXCEL_EXPORT_DIR, filename))

    df = build_dataframe(codes)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="QR Codes")
        ws = writer.sheets["QR Codes"]

        # Header styling
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        for col_idx, _col in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Auto-width columns
        for col_idx, col in enumerate(COLUMNS, start=1):
            max_len = max((len(str(v)) for v in df[col].astype(str)), default=10)
            max_len = max(max_len, len(col))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

        ws.freeze_panes = "A2"

    logger.info("Excel export written: %s (%d rows)", path, len(df))
    return path
